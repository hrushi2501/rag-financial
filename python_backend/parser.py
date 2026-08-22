"""Page-preserving document extraction.

PyMuPDF4LLM is used when installed and falls back to pypdf. The fallback keeps
the backend usable in lightweight development and test environments.
"""

from __future__ import annotations

import csv
import os
import tempfile
import zipfile
from io import BytesIO, StringIO
from pathlib import Path

try:
    from .models import PageContent
except ImportError:  # pragma: no cover
    from models import PageContent


class FinancialDocumentParser:
    def __init__(self, minimum_pdf_text: int = 50, max_pages: int = 2000, max_archive_size: int = 536870912):
        self.minimum_pdf_text = minimum_pdf_text
        self.max_pages = max_pages
        self.max_archive_size = max_archive_size

    def parse_path(self, path: str | Path, filename: str | None = None) -> list[PageContent]:
        file_path = Path(path)
        return self.parse_bytes(filename or file_path.name, file_path.read_bytes())

    def parse_bytes(self, filename: str, content: bytes) -> list[PageContent]:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf(content)
        if ext == ".docx":
            self._validate_archive(content)
            return self._parse_docx(content)
        if ext == ".xlsx":
            self._validate_archive(content)
            return self._parse_xlsx(content)
        if ext == ".csv":
            return [PageContent(1, self._parse_csv(content))]
        if ext == ".txt":
            return [PageContent(1, content.decode("utf-8", errors="replace"))]
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")

    def _parse_pdf(self, content: bytes) -> list[PageContent]:
        try:
            return self._parse_pdf_markdown(content)
        except ImportError:
            return self._parse_pdf_fallback(content)
        except Exception:
            # Parser-specific failures should not make native-text PDFs unusable.
            return self._parse_pdf_fallback(content)

    def _parse_pdf_markdown(self, content: bytes) -> list[PageContent]:
        import pymupdf
        import pymupdf4llm

        temporary_name = ""
        try:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as temporary:
                temporary.write(content)
                temporary_name = temporary.name
            document = pymupdf.open(temporary_name)
            try:
                if len(document) > self.max_pages:
                    raise ValueError(f"Document exceeds the {self.max_pages}-page safety limit")
                items = pymupdf4llm.to_markdown(
                    document,
                    page_chunks=True,
                    write_images=False,
                    table_strategy="lines_strict",
                    show_progress=False,
                )
            finally:
                document.close()
        finally:
            if temporary_name:
                try:
                    os.unlink(temporary_name)
                except OSError:
                    pass

        pages = []
        for index, item in enumerate(items):
            metadata = item.get("metadata") or {}
            raw_page = metadata.get("page", index)
            # PyMuPDF4LLM versions differ between zero and one based metadata.
            page_number = int(raw_page) + 1 if int(raw_page) == index else int(raw_page)
            text = (item.get("text") or "").strip()
            tables = tuple(str(table) for table in item.get("tables", []) if table)
            pages.append(PageContent(max(1, page_number), text, tables))
        self._validate_pdf(pages)
        return pages

    def _parse_pdf_fallback(self, content: bytes) -> list[PageContent]:
        import pypdf

        reader = pypdf.PdfReader(BytesIO(content))
        if len(reader.pages) > self.max_pages:
            raise ValueError(f"Document exceeds the {self.max_pages}-page safety limit")
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ValueError("PDF is encrypted") from exc
        pages = [
            PageContent(page_number=index + 1, text=(page.extract_text() or "").strip())
            for index, page in enumerate(reader.pages)
        ]
        self._validate_pdf(pages)
        return pages

    def _validate_pdf(self, pages: list[PageContent]) -> None:
        if sum(len(page.text) for page in pages) < self.minimum_pdf_text:
            raise ValueError("PDF contains insufficient native text; OCR is required")

    def _validate_archive(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 20_000:
                    raise ValueError("Office document contains too many archive entries")
                if sum(member.file_size for member in members) > self.max_archive_size:
                    raise ValueError("Office document expands beyond the safety limit")
        except zipfile.BadZipFile as exc:
            raise ValueError("Office document is not a valid ZIP container") from exc

    @staticmethod
    def _parse_docx(content: bytes) -> list[PageContent]:
        from docx import Document

        document = Document(BytesIO(content))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            rows = [" | ".join(str(cell.text).strip() for cell in row.cells) for row in table.rows]
            if rows:
                parts.append("\n".join(rows))
        return [PageContent(1, "\n\n".join(parts))]

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[PageContent]:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        pages = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows = [
                " | ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            pages.append(PageContent(index, f"# Sheet: {sheet.title}\n\n" + "\n".join(filter(str.strip, rows))))
        workbook.close()
        return pages

    @staticmethod
    def _parse_csv(content: bytes) -> str:
        reader = csv.reader(StringIO(content.decode("utf-8-sig", errors="replace")))
        return "\n".join(" | ".join(row) for row in reader)
